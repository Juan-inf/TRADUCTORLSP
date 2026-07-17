"""
generar_docx_plan_despliegue.py — Convierte ENTREGABLE_PLAN_DE_DESPLIEGUE_S13.md → .docx
  - Portada con metadatos del entregable
  - Estilos: Normal, Heading 1-3, Code, Table
  - Tablas y bloques de código Markdown renderizados
"""

import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).parent.parent
MD_IN = ROOT / "ENTREGABLE_PLAN_DE_DESPLIEGUE_S13.md"
DOCX_OUT = ROOT / "ENTREGABLE_PLAN_DE_DESPLIEGUE_S13.docx"


def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


print("Creando ENTREGABLE_PLAN_DE_DESPLIEGUE_S13.docx...")
doc = Document()

for section in doc.sections:
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.5)

style_normal = doc.styles["Normal"]
style_normal.font.name = "Calibri"
style_normal.font.size = Pt(11)
style_normal.paragraph_format.space_after = Pt(6)

h1 = doc.styles["Heading 1"]
h1.font.name = "Calibri"; h1.font.size = Pt(16); h1.font.bold = True
h1.font.color.rgb = RGBColor(0x0e, 0x7c, 0x78)
h1.paragraph_format.space_before = Pt(18); h1.paragraph_format.space_after = Pt(6)

h2 = doc.styles["Heading 2"]
h2.font.name = "Calibri"; h2.font.size = Pt(13); h2.font.bold = True
h2.font.color.rgb = RGBColor(0x0e, 0x7c, 0x78)
h2.paragraph_format.space_before = Pt(12); h2.paragraph_format.space_after = Pt(4)

h3 = doc.styles["Heading 3"]
h3.font.name = "Calibri"; h3.font.size = Pt(11); h3.font.bold = True
h3.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
h3.paragraph_format.space_before = Pt(10); h3.paragraph_format.space_after = Pt(3)

if "Code" not in [s.name for s in doc.styles]:
    code_style = doc.styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
else:
    code_style = doc.styles["Code"]
code_style.font.name = "Courier New"
code_style.font.size = Pt(9)
code_style.paragraph_format.space_before = Pt(4)
code_style.paragraph_format.space_after = Pt(4)
code_style.paragraph_format.left_indent = Cm(0.5)

# ── PORTADA ──────────────────────────────────────────────────────────────────
doc.add_paragraph()

title_para = doc.add_paragraph()
title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title_para.add_run("PLAN DE DESPLIEGUE")
run.font.name = "Calibri"; run.font.size = Pt(22); run.font.bold = True
run.font.color.rgb = RGBColor(0x0e, 0x7c, 0x78)

sub_para = doc.add_paragraph()
sub_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub_run = sub_para.add_run("Traductor LSP → Castellano\nEntregable S13 — antes de la tesis")
sub_run.font.name = "Calibri"; sub_run.font.size = Pt(13); sub_run.font.italic = True
sub_run.font.color.rgb = RGBColor(0x42, 0x50, 0x55)

doc.add_paragraph()
meta_table = doc.add_table(rows=4, cols=2)
meta_table.style = "Table Grid"
meta_data = [
    ("Proyecto", "Traductor LSP → Castellano"),
    ("Fecha", "2026-07-17"),
    ("Responsable", "Juan Calla"),
    ("Estado", "6/6 tests pasan · E2E verificado en vivo"),
]
for i, (key, val) in enumerate(meta_data):
    row = meta_table.rows[i]
    row.cells[0].text = key
    row.cells[1].text = val
    set_cell_bg(row.cells[0], "0e7c78")
    run0 = row.cells[0].paragraphs[0].runs[0]
    run0.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    run0.font.bold = True
    run0.font.size = Pt(10)
    row.cells[1].paragraphs[0].runs[0].font.size = Pt(10)

doc.add_page_break()

# ── Parser Markdown → DOCX ────────────────────────────────────────────────────
def strip_inline(text):
    text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
    text = re.sub(r'\*(.+?)\*', r'\1', text)
    text = re.sub(r'`(.+?)`', r'\1', text)
    return text


def add_rich_run(para, text):
    pattern = r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)'
    parts = re.split(pattern, text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = para.add_run(part[2:-2]); r.bold = True
        elif part.startswith("*") and part.endswith("*"):
            r = para.add_run(part[1:-1]); r.italic = True
        elif part.startswith("`") and part.endswith("`"):
            r = para.add_run(part[1:-1]); r.font.name = "Courier New"; r.font.size = Pt(9)
        else:
            para.add_run(part)


raw = MD_IN.read_text(encoding="utf-8")
lines = raw.splitlines()

# Saltar portada MD (título + metadatos) hasta la primera sección "## 1."
i = 0
while i < len(lines) and not lines[i].startswith("## 1."):
    i += 1

total = len(lines)
in_code = False
in_table = False
code_lines = []
table_lines = []

while i < total:
    line = lines[i]

    if line.startswith("```"):
        if not in_code:
            in_code = True
            code_lines = []
        else:
            in_code = False
            for cl in "\n".join(code_lines).split("\n"):
                p = doc.add_paragraph(cl, style="Code")
                pPr = p._p.get_or_add_pPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "F1F5F9")
                pPr.append(shd)
        i += 1; continue

    if in_code:
        code_lines.append(line); i += 1; continue

    if line.startswith("|"):
        if not in_table:
            in_table = True; table_lines = []
        table_lines.append(line); i += 1; continue
    else:
        if in_table:
            in_table = False
            rows = [r for r in table_lines if not re.match(r'^\|[\s\-:|]+\|', r)]
            if rows:
                n_cols = rows[0].count("|") - 1
                tbl = doc.add_table(rows=len(rows), cols=max(n_cols, 1))
                tbl.style = "Table Grid"
                for ri, row_line in enumerate(rows):
                    cells = [c.strip() for c in row_line.strip("|").split("|")]
                    for ci, cell_text in enumerate(cells):
                        if ci < len(tbl.rows[ri].cells):
                            cell = tbl.rows[ri].cells[ci]
                            cell.paragraphs[0].clear()
                            add_rich_run(cell.paragraphs[0], strip_inline(cell_text))
                            if cell.paragraphs[0].runs:
                                cell.paragraphs[0].runs[0].font.size = Pt(9)
                            if ri == 0:
                                set_cell_bg(cell, "0e7c78")
                                for run in cell.paragraphs[0].runs:
                                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                                    run.font.bold = True
                doc.add_paragraph()

    h3_match = re.match(r'^### (.+)', line)
    h2_match = re.match(r'^## (.+)', line)
    h1_match = re.match(r'^# (.+)', line)

    if h1_match:
        doc.add_heading(strip_inline(h1_match.group(1)), level=1); i += 1; continue
    if h2_match:
        doc.add_heading(strip_inline(h2_match.group(1)), level=2); i += 1; continue
    if h3_match:
        doc.add_heading(strip_inline(h3_match.group(1)), level=3); i += 1; continue

    if re.match(r'^---+$', line.strip()):
        i += 1; continue

    bullet_match = re.match(r'^(\s*)[-*] (.+)', line)
    if bullet_match:
        indent = len(bullet_match.group(1)) // 2
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(indent * 0.5)
        add_rich_run(p, bullet_match.group(2))
        i += 1; continue

    num_match = re.match(r'^\d+\. (.+)', line)
    if num_match:
        p = doc.add_paragraph(style="List Number")
        add_rich_run(p, num_match.group(1))
        i += 1; continue

    bq_match = re.match(r'^> (.+)', line)
    if bq_match:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8)
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single"); left.set(qn("w:sz"), "12")
        left.set(qn("w:space"), "12"); left.set(qn("w:color"), "0e7c78")
        pBdr.append(left); pPr.append(pBdr)
        add_rich_run(p, bq_match.group(1))
        i += 1; continue

    if not line.strip():
        i += 1; continue

    p = doc.add_paragraph()
    add_rich_run(p, line)
    i += 1

print(f"Guardando {DOCX_OUT.name}...")
doc.save(str(DOCX_OUT))
size_kb = DOCX_OUT.stat().st_size / 1024
print(f"OK — {DOCX_OUT.name} ({size_kb:.0f} KB) en {DOCX_OUT}")
