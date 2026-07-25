"""generar_docx_resultados_presentacion.py — RESULTADOS_PRESENTACION_S15.md → .docx

Portada + secciones con tablas, figuras embebidas y texto formateado.
Adaptado de generar_docx_resultados_sustentacion.py para el documento
consolidado de resultados (ensemble v4+S29, activo desde 2026-07-19).
"""
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).parent.parent
MD_IN = ROOT / "RESULTADOS_PRESENTACION_S15.md"
DOCX_OUT = ROOT / "RESULTADOS_PRESENTACION_S15.docx"


def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


print("Creando RESULTADOS_PRESENTACION_S15.docx...")
doc = Document()

for section in doc.sections:
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.5)

style_normal = doc.styles["Normal"]
style_normal.font.name = "Calibri"
style_normal.font.size = Pt(11)
style_normal.paragraph_format.space_after = Pt(6)

h1 = doc.styles["Heading 1"]
h1.font.name = "Calibri"; h1.font.size = Pt(16); h1.font.bold = True
h1.font.color.rgb = RGBColor(0x0e, 0x7c, 0x78)
h1.paragraph_format.space_before = Pt(18); h1.paragraph_format.space_after = Pt(6)

h2 = doc.styles["Heading 2"]
h2.font.name = "Calibri"; h2.font.size = Pt(13); h2.font.bold = True
h2.font.color.rgb = RGBColor(0x0e, 0x7c, 0x78)
h2.paragraph_format.space_before = Pt(12); h2.paragraph_format.space_after = Pt(4)

h3 = doc.styles["Heading 3"]
h3.font.name = "Calibri"; h3.font.size = Pt(11); h3.font.bold = True
h3.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
h3.paragraph_format.space_before = Pt(10); h3.paragraph_format.space_after = Pt(3)

if "Caption" in [s.name for s in doc.styles]:
    cap_style = doc.styles["Caption"]
    cap_style.font.name = "Calibri"; cap_style.font.size = Pt(9)
    cap_style.font.italic = True
    cap_style.font.color.rgb = RGBColor(0x64, 0x74, 0x8b)

# ── PORTADA ──────────────────────────────────────────────────────────────────
doc.add_paragraph()
title_para = doc.add_paragraph()
title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title_para.add_run("RESULTADOS")
run.font.name = "Calibri"; run.font.size = Pt(24); run.font.bold = True
run.font.color.rgb = RGBColor(0x0e, 0x7c, 0x78)

sub_para = doc.add_paragraph()
sub_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub_run = sub_para.add_run("Traductor LSP → Castellano\nDocumento consolidado para presentación y sustentación de tesis")
sub_run.font.name = "Calibri"; sub_run.font.size = Pt(13); sub_run.font.italic = True
sub_run.font.color.rgb = RGBColor(0x42, 0x50, 0x55)

doc.add_paragraph()
meta_table = doc.add_table(rows=4, cols=2)
meta_table.style = "Table Grid"
meta_data = [
    ("Proyecto", "Traductor LSP → Castellano"),
    ("Fecha de consolidación", "2026-07-23"),
    ("Autor", "Juan Calla"),
    ("Modelo activo", "Ensemble BiLSTM S27(v4) + S29 — F1-macro=0.4424 (holdout limpio)"),
]
for i, (key, val) in enumerate(meta_data):
    row = meta_table.rows[i]
    row.cells[0].text = key
    row.cells[1].text = val
    set_cell_bg(row.cells[0], "0e7c78")
    run0 = row.cells[0].paragraphs[0].runs[0]
    run0.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    run0.font.bold = True
    run0.font.size = Pt(10)
    row.cells[1].paragraphs[0].runs[0].font.size = Pt(10)

doc.add_page_break()


def strip_inline(text):
    text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
    text = re.sub(r'\*(.+?)\*', r'\1', text)
    text = re.sub(r'`(.+?)`', r'\1', text)
    return text


def add_rich_run(para, text):
    pattern = r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)'
    parts = re.split(pattern, text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = para.add_run(part[2:-2]); r.bold = True
        elif part.startswith("*") and part.endswith("*"):
            r = para.add_run(part[1:-1]); r.italic = True
        elif part.startswith("`") and part.endswith("`"):
            r = para.add_run(part[1:-1]); r.font.name = "Courier New"; r.font.size = Pt(9)
        else:
            para.add_run(part)


raw = MD_IN.read_text(encoding="utf-8")
lines = raw.splitlines()

# Saltar portada del markdown (título, subtítulo, línea de nota) hasta la 1a sección "## 1."
i = 0
while i < len(lines) and not lines[i].startswith("## 1."):
    i += 1

total = len(lines)
in_table = False
table_lines = []

while i < total:
    line = lines[i]

    if line.startswith("|"):
        if not in_table:
            in_table = True; table_lines = []
        table_lines.append(line); i += 1; continue
    else:
        if in_table:
            in_table = False
            rows = [r for r in table_lines if not re.match(r'^\|[\s\-:|]+\|', r)]
            if rows:
                n_cols = rows[0].count("|") - 1
                tbl = doc.add_table(rows=len(rows), cols=max(n_cols, 1))
                tbl.style = "Table Grid"
                for ri, row_line in enumerate(rows):
                    cells = [c.strip() for c in row_line.strip("|").split("|")]
                    for ci, cell_text in enumerate(cells):
                        if ci < len(tbl.rows[ri].cells):
                            cell = tbl.rows[ri].cells[ci]
                            cell.paragraphs[0].clear()
                            add_rich_run(cell.paragraphs[0], strip_inline(cell_text))
                            if cell.paragraphs[0].runs:
                                cell.paragraphs[0].runs[0].font.size = Pt(9)
                            if ri == 0:
                                set_cell_bg(cell, "0e7c78")
                                for run in cell.paragraphs[0].runs:
                                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                                    run.font.bold = True
                doc.add_paragraph()

    img_match = re.match(r'^!\[([^\]]*)\]\(([^)]+)\)', line)
    if img_match:
        alt, path = img_match.group(1), img_match.group(2)
        img_path = ROOT / path
        if img_path.exists():
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run()
            run.add_picture(str(img_path), width=Inches(5.8))
            cap = doc.add_paragraph(alt)
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if cap.runs:
                cap.runs[0].font.size = Pt(9)
                cap.runs[0].font.italic = True
                cap.runs[0].font.color.rgb = RGBColor(0x64, 0x74, 0x8b)
        i += 1; continue

    h3_match = re.match(r'^### (.+)', line)
    h2_match = re.match(r'^## (.+)', line)
    h1_match = re.match(r'^# (.+)', line)
    if h1_match:
        doc.add_heading(strip_inline(h1_match.group(1)), level=1); i += 1; continue
    if h2_match:
        doc.add_heading(strip_inline(h2_match.group(1)), level=2); i += 1; continue
    if h3_match:
        doc.add_heading(strip_inline(h3_match.group(1)), level=3); i += 1; continue

    if re.match(r'^---+$', line.strip()):
        i += 1; continue

    if re.match(r'^> ', line):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8)
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single"); left.set(qn("w:sz"), "12")
        left.set(qn("w:space"), "12"); left.set(qn("w:color"), "0e7c78")
        pBdr.append(left); pPr.append(pBdr)
        add_rich_run(p, line[2:])
        i += 1; continue

    num_match = re.match(r'^\d+\. (.+)', line)
    if num_match:
        p = doc.add_paragraph(style="List Number")
        add_rich_run(p, num_match.group(1))
        i += 1; continue

    bullet_match = re.match(r'^[-*] (.+)', line)
    if bullet_match:
        p = doc.add_paragraph(style="List Bullet")
        add_rich_run(p, bullet_match.group(1))
        i += 1; continue

    if not line.strip():
        i += 1; continue

    line = re.sub(r'!\[([^\]]*)\]\([^)]+\)', r'(ver figura: \1)', line)
    p = doc.add_paragraph()
    add_rich_run(p, line)
    i += 1

print(f"Guardando {DOCX_OUT.name}...")
doc.save(str(DOCX_OUT))
size_kb = DOCX_OUT.stat().st_size / 1024
print(f"OK — {DOCX_OUT.name} ({size_kb:.0f} KB) en {DOCX_OUT}")
